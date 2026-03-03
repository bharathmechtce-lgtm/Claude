#!/usr/bin/env python3
"""
Generate an Excel workbook with all failed scenarios for manual testing.

Sheets:
  1. Summary       — one row per failure with key metrics
  2. Per-failure    — one sheet per failure with full messages + SAP truth + previous results

Usage:
    python testing/scenarios/build_failure_excel.py
"""

import json
import os
import sys
from datetime import datetime

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("pip install openpyxl")
    sys.exit(1)

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
BENCHMARKS_DIR = os.path.join(os.path.dirname(SCRIPT_DIR), "benchmarks")
RESULTS_DIR = os.path.join(os.path.dirname(SCRIPT_DIR), "results")

# ── Styles ──
HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
PASS_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
FAIL_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
BOTH_FAIL_FILL = PatternFill(start_color="FF9999", end_color="FF9999", fill_type="solid")
MSG_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
SAP_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
DETAIL_FONT = Font(name="Calibri", size=10)
BOLD_FONT = Font(name="Calibri", bold=True, size=10)
WRAP = Alignment(wrap_text=True, vertical="top")
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)


def style_header_row(ws, row, max_col):
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER


def auto_width(ws, min_width=10, max_width=60):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        lengths = []
        for cell in col:
            if cell.value:
                lines = str(cell.value).split("\n")
                lengths.append(max(len(l) for l in lines))
        if lengths:
            width = min(max(max(lengths) + 2, min_width), max_width)
            ws.column_dimensions[letter].width = width


def load_data():
    with open(os.path.join(BENCHMARKS_DIR, "test_scenarios.json")) as f:
        scenarios = json.load(f)["scenarios"]
    scenarios_by_id = {}
    for s in scenarios:
        scenarios_by_id[f"S{s['scenario_id']:02d}"] = s

    comparison_files = sorted(
        [f for f in os.listdir(RESULTS_DIR) if f.startswith("comparison_") and f.endswith(".json")],
        reverse=True,
    )
    if not comparison_files:
        print("No comparison results found in testing/results/")
        sys.exit(1)

    with open(os.path.join(RESULTS_DIR, comparison_files[0])) as f:
        results = json.load(f)

    return scenarios_by_id, results


def identify_failures(results):
    """Return list of (scenario_id, ship_to, haiku_result, gemini_result) for all failures."""
    haiku = {(r["scenario_id"], r["ship_to"]): r for r in results.get("claude-haiku-4-5-20251001", [])}
    gemini = {(r["scenario_id"], r["ship_to"]): r for r in results.get("gemini-2.5-flash", [])}

    failures = []
    all_keys = set(haiku.keys()) | set(gemini.keys())
    for key in sorted(all_keys):
        h = haiku.get(key, {})
        g = gemini.get(key, {})
        if not h.get("complete_order", False) or not g.get("complete_order", False):
            failures.append((key[0], key[1], h, g))
    return failures


def build_summary_sheet(wb, failures, scenarios_by_id):
    ws = wb.active
    ws.title = "Summary"

    headers = [
        "Failure #", "Scenario", "Ship To", "Difficulty", "Chat Name",
        "Target Items", "Customer Messages",
        "Haiku Score", "Haiku Status", "Gemini Score", "Gemini Status",
        "Who Failed", "Root Cause Pattern",
    ]
    for ci, h in enumerate(headers, 1):
        ws.cell(row=1, column=ci, value=h)
    style_header_row(ws, 1, len(headers))

    for i, (sid, ship_to, h, g) in enumerate(failures, 1):
        scenario = scenarios_by_id.get(sid, {})

        # Collect customer messages for this scenario
        msgs = []
        for conv in scenario.get("conversations_1to1", []):
            for m in conv.get("messages", []):
                if m.get("type") in ("order", "order_addition"):
                    msgs.append(m.get("text", "")[:100])

        h_complete = h.get("complete_order", False)
        g_complete = g.get("complete_order", False)
        h_score = h.get("order_score", 0)
        g_score = g.get("order_score", 0)

        if not h_complete and not g_complete:
            who = "BOTH"
        elif not h_complete:
            who = "Haiku"
        else:
            who = "Gemini"

        # Root cause pattern
        h_extra = h.get("extra", 0)
        g_extra = g.get("extra", 0)
        g_count = g.get("llm_count", 0)
        h_count = h.get("llm_count", 0)

        if g_count == 0 and not g_complete:
            pattern = "Gemini zero output"
        elif h_extra >= 3 and not h_complete:
            pattern = "Haiku multi-outlet extras"
        elif not h.get("ship_to_correct", True):
            pattern = "Ship-to name mismatch"
        elif who == "BOTH":
            pattern = "Both models struggled"
        else:
            pattern = "Item/qty mismatch"

        row = [
            i, sid, ship_to, scenario.get("difficulty", ""),
            scenario.get("chat_name", ""),
            h.get("target_count", g.get("target_count", 0)),
            "\n".join(msgs),
            f"{h_score:.0f}%", "PASS" if h_complete else "FAIL",
            f"{g_score:.0f}%", "PASS" if g_complete else "FAIL",
            who, pattern,
        ]

        for ci, v in enumerate(row, 1):
            cell = ws.cell(row=i + 1, column=ci, value=v)
            cell.font = DETAIL_FONT
            cell.alignment = WRAP
            cell.border = THIN_BORDER

        # Color the status cells
        h_cell = ws.cell(row=i + 1, column=9)
        g_cell = ws.cell(row=i + 1, column=11)
        h_cell.fill = PASS_FILL if h_complete else FAIL_FILL
        g_cell.fill = PASS_FILL if g_complete else FAIL_FILL

        who_cell = ws.cell(row=i + 1, column=12)
        if who == "BOTH":
            who_cell.fill = BOTH_FAIL_FILL

    auto_width(ws)
    ws.freeze_panes = "A2"


def build_detail_sheets(wb, failures, scenarios_by_id):
    for i, (sid, ship_to, h, g) in enumerate(failures, 1):
        scenario = scenarios_by_id.get(sid, {})
        # Sheet name max 31 chars
        sheet_name = f"F{i:02d}_{sid}_{ship_to[:18]}"[:31]
        ws = wb.create_sheet(title=sheet_name)

        row = 1

        # ── Header ──
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        cell = ws.cell(row=row, column=1,
                       value=f"FAILURE #{i}: {sid} | {ship_to}")
        cell.font = Font(name="Calibri", bold=True, size=14, color="2F5496")
        row += 1

        # ── Metadata ──
        meta = [
            ("Difficulty", scenario.get("difficulty", "")),
            ("Chat Name", scenario.get("chat_name", "")),
            ("Date", scenario.get("date_iso", "")),
            ("Ship To", ship_to),
            ("Target Items", str(h.get("target_count", g.get("target_count", 0)))),
            ("Haiku Score", f"{h.get('order_score', 0):.0f}% ({'PASS' if h.get('complete_order') else 'FAIL'})"),
            ("Gemini Score", f"{g.get('order_score', 0):.0f}% ({'PASS' if g.get('complete_order') else 'FAIL'})"),
        ]
        for label, val in meta:
            ws.cell(row=row, column=1, value=label).font = BOLD_FONT
            ws.cell(row=row, column=2, value=val).font = DETAIL_FONT
            row += 1

        row += 1

        # ── Customer Messages (ALL from this scenario) ──
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        cell = ws.cell(row=row, column=1, value="CUSTOMER MESSAGES (from WhatsApp)")
        cell.fill = MSG_FILL
        cell.font = Font(name="Calibri", bold=True, size=11)
        row += 1

        headers = ["Time", "Sender", "Type", "Message Text"]
        for ci, hdr in enumerate(headers, 1):
            c = ws.cell(row=row, column=ci, value=hdr)
            c.fill = MSG_FILL
            c.font = BOLD_FONT
            c.border = THIN_BORDER
        row += 1

        for conv in scenario.get("conversations_1to1", []):
            for m in conv.get("messages", []):
                ws.cell(row=row, column=1, value=m.get("time", "")).font = DETAIL_FONT
                ws.cell(row=row, column=2, value=m.get("sender", "")).font = DETAIL_FONT
                ws.cell(row=row, column=3, value=m.get("type", "")).font = DETAIL_FONT
                text_cell = ws.cell(row=row, column=4, value=m.get("text", ""))
                text_cell.font = DETAIL_FONT
                text_cell.alignment = WRAP
                for ci in range(1, 5):
                    ws.cell(row=row, column=ci).border = THIN_BORDER
                row += 1

        row += 1

        # ── SAP Expected Order (for THIS ship_to) ──
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        cell = ws.cell(row=row, column=1, value=f"SAP EXPECTED ORDER (for {ship_to})")
        cell.fill = SAP_FILL
        cell.font = Font(name="Calibri", bold=True, size=11)
        row += 1

        sap_headers = ["#", "Item Code", "Description", "Quantity", "Price", "Doc Num"]
        for ci, hdr in enumerate(sap_headers, 1):
            c = ws.cell(row=row, column=ci, value=hdr)
            c.fill = SAP_FILL
            c.font = BOLD_FONT
            c.border = THIN_BORDER
        row += 1

        sap_items = [t for t in scenario.get("sap_truth", [])
                     if t.get("ship_to_code", "") == ship_to]
        for idx, item in enumerate(sap_items, 1):
            ws.cell(row=row, column=1, value=idx).font = DETAIL_FONT
            ws.cell(row=row, column=2, value=item.get("item_code", "")).font = DETAIL_FONT
            ws.cell(row=row, column=3, value=item.get("description", "")).font = DETAIL_FONT
            ws.cell(row=row, column=4, value=item.get("quantity", 0)).font = DETAIL_FONT
            ws.cell(row=row, column=5, value=item.get("price", 0)).font = DETAIL_FONT
            ws.cell(row=row, column=6, value=item.get("doc_num", "")).font = DETAIL_FONT
            for ci in range(1, 7):
                ws.cell(row=row, column=ci).border = THIN_BORDER
            row += 1

        row += 1

        # ── Previous Results Detail ──
        for model_label, result in [("HAIKU", h), ("GEMINI", g)]:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
            status = "PASS" if result.get("complete_order") else "FAIL"
            cell = ws.cell(row=row, column=1,
                           value=f"PREVIOUS {model_label} RESULT ({status} — {result.get('order_score', 0):.0f}%)")
            cell.font = Font(name="Calibri", bold=True, size=11,
                             color="006100" if status == "PASS" else "9C0006")
            row += 1

            detail_headers = ["Status", "LLM Item", "LLM Code", "LLM Qty", "SAP Item", "SAP Code", "SAP Qty"]
            for ci, hdr in enumerate(detail_headers, 1):
                c = ws.cell(row=row, column=ci, value=hdr)
                c.font = BOLD_FONT
                c.border = THIN_BORDER
            row += 1

            for d in result.get("details", []):
                ws.cell(row=row, column=1, value=d.get("status", "")).font = DETAIL_FONT
                ws.cell(row=row, column=2, value=d.get("llm_item", "")).font = DETAIL_FONT
                ws.cell(row=row, column=3, value=d.get("llm_code", "")).font = DETAIL_FONT
                ws.cell(row=row, column=4, value=d.get("llm_qty", 0)).font = DETAIL_FONT
                ws.cell(row=row, column=5, value=d.get("sap_item", "")).font = DETAIL_FONT
                ws.cell(row=row, column=6, value=d.get("sap_code", "")).font = DETAIL_FONT
                ws.cell(row=row, column=7, value=d.get("sap_qty", 0)).font = DETAIL_FONT

                status_cell = ws.cell(row=row, column=1)
                if d.get("status") == "MATCH":
                    status_cell.fill = PASS_FILL
                elif d.get("status") in ("MISSED", "EXTRA", "QTY_MISMATCH"):
                    status_cell.fill = FAIL_FILL

                for ci in range(1, 8):
                    ws.cell(row=row, column=ci).border = THIN_BORDER
                row += 1

            if not result.get("details"):
                ws.cell(row=row, column=1, value="(No output — 0 items produced)").font = DETAIL_FONT
                row += 1

            row += 1

        auto_width(ws)


def main():
    print("Loading data...")
    scenarios_by_id, results = load_data()
    failures = identify_failures(results)
    print(f"Found {len(failures)} failures")

    wb = Workbook()
    build_summary_sheet(wb, failures, scenarios_by_id)
    build_detail_sheets(wb, failures, scenarios_by_id)

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_path = os.path.join(RESULTS_DIR, f"failed_scenarios_{ts}.xlsx")
    wb.save(out_path)
    print(f"\nSaved: {out_path}")
    print(f"  - Summary sheet: {len(failures)} rows")
    print(f"  - Detail sheets: {len(failures)} sheets (one per failure)")


if __name__ == "__main__":
    main()
